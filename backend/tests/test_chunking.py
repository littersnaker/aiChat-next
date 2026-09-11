"""知识库文档切块与格式解析测试。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from backend.services.embeddings.chunking import (
    extract_document_text,
    split_text,
)


def test_split_text_returns_empty_for_blank() -> None:
    """空白文本直接返回空列表。"""

    assert split_text("", max_chars=100, overlap=10) == []
    assert split_text("   \n  ", max_chars=100, overlap=10) == []


def test_split_text_short_text_single_piece() -> None:
    """不超过上限的文本整块返回。"""

    assert split_text("你好，世界", max_chars=100, overlap=10) == ["你好，世界"]


def test_split_text_respects_max_chars_and_overlap() -> None:
    """长文本切块后每块不超过上限，且相邻块有重叠。"""

    text = "".join(f"第{i}句。" for i in range(60))
    pieces = split_text(text, max_chars=50, overlap=10)
    assert len(pieces) > 1
    assert all(len(piece) <= 50 for piece in pieces)
    # 除最后一块外，其余块应达到上限（保证块间重叠存在）。
    assert all(len(piece) == 50 for piece in pieces[:-1])
    # 重叠：后一块的开头应出现在前一块的尾部。
    assert pieces[1][:10] in pieces[0]


def test_extract_markdown_and_text(tmp_path: Path) -> None:
    """markdown/txt 直接读取原文。"""

    markdown = tmp_path / "a.md"
    markdown.write_text("# 标题\n正文", encoding="utf-8")
    assert extract_document_text(markdown) == "# 标题\n正文"

    markdown2 = tmp_path / "b.markdown"
    markdown2.write_text("## 二级\n内容", encoding="utf-8")
    assert extract_document_text(markdown2) == "## 二级\n内容"

    plain = tmp_path / "c.txt"
    plain.write_text("纯文本内容", encoding="utf-8")
    assert extract_document_text(plain) == "纯文本内容"


def test_extract_csv_including_bom(tmp_path: Path) -> None:
    """CSV 解析按行用 | 连接，并兼容 Excel 导出的 BOM 文件。"""

    csv_path = tmp_path / "d.csv"
    csv_path.write_text("城市,温度\n上海,28\n", encoding="utf-8-sig")
    text = extract_document_text(csv_path)
    assert text == "城市 | 温度\n上海 | 28"


def test_extract_xlsx_reads_all_sheets(tmp_path: Path) -> None:
    """XLSX 逐工作表抽取，行内 | 连接，带工作表标题前缀。"""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "货架"
    sheet.append(["货位", "数量"])
    sheet.append(["A-01", "24"])
    sheet2 = workbook.create_sheet("备注")
    sheet2.append([" onlySecond "])
    xlsx_path = tmp_path / "e.xlsx"
    workbook.save(xlsx_path)

    text = extract_document_text(xlsx_path)
    assert "[货架] 货位 | 数量" in text
    assert "[货架] A-01 | 24" in text
    assert "[备注] onlySecond" in text


def test_extract_xlsx_ignores_blank_rows(tmp_path: Path) -> None:
    """全空行不产生空串行。"""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "S"
    sheet.append(["a", None, "b"])
    sheet.append([None, None, None])
    xlsx_path = tmp_path / "f.xlsx"
    workbook.save(xlsx_path)

    text = extract_document_text(xlsx_path)
    assert text == "[S] a | b"


def test_extract_unsupported_suffix_returns_empty(tmp_path: Path) -> None:
    """未知扩展名返回空串而不是抛异常。"""

    binary = tmp_path / "g.exe"
    binary.write_bytes(b"MZ...")
    assert extract_document_text(binary) == ""
